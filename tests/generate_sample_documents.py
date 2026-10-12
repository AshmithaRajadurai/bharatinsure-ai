"""
Generate rich, multi-page sample PDF, DOCX, and TXT insurance policies for automated testing.
Includes realistic policy clauses covering:
- Policy Schedule (Start, expiry dates, grace period, co-payment overrides)
- Waiting periods (PED 36m, initial 30d, specific diseases 24m, maternity 24m add-on)
- In-patient hospitalisation (Room rent zero capping, ICU, AYUSH, Day care)
- Riders & modern treatments (Air ambulance INR 2,50,000, Robotic surgery 10% copay, restoration)
- Permanent exclusions (Cosmetic, bariatric unless BMI > 40, dental trauma exception)
- Claim procedures (48h planned, 24h emergency, 30d bills submission, rejection grounds)
"""

import os
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
import docx

os.makedirs("tests/sample_docs", exist_ok=True)

# ---------------------------------------------------------
# 1. Generate Comprehensive 6-Page HDFC Optima Secure PDF
# ---------------------------------------------------------
pdf_path = "tests/sample_docs/HDFC_Optima_Secure_Policy.pdf"
c = canvas.Canvas(pdf_path, pagesize=letter)

# Page 1: Policy Schedule & Identification
c.setFont("Helvetica-Bold", 16)
c.drawString(50, 750, "HDFC ERGO General Insurance Company Limited")
c.setFont("Helvetica-Bold", 13)
c.drawString(50, 725, "Optima Secure Individual Health Insurance Policy Document")
c.setFont("Helvetica", 11)
c.drawString(50, 695, "Policy Number: HDFC-HEALTH-2026-789012")
c.drawString(50, 675, "Policyholder: Rajesh Sharma | Sum Insured: INR 10,00,000 (Ten Lakhs)")
c.drawString(50, 655, "Policy Start Date: 01 April 2026 | Policy Expiry Date: 31 March 2027")
c.drawString(50, 635, "Renewal Grace Period: 30 days from expiry date for uninterrupted continuity benefits.")
c.drawString(50, 605, "Section 1: Policy Schedule & Co-payment Terms")
c.drawString(50, 585, "1.1 Standard Co-payment: Zero percent (0%) across all network hospitals for general claims.")
c.drawString(50, 565, "1.2 Special Co-payment Override: As noted in Schedule, Section 4.2 robotic surgery")
c.drawString(50, 545, "requires a mandatory 10% co-payment notwithstanding Section 1.1.")
c.drawString(50, 515, "1.3 Cumulative Bonus: 50% increase in basic Sum Insured for every claim-free year, up to 100%.")
c.drawString(50, 40, "Page 1 of 6")
c.showPage()

# Page 2: Waiting Periods & PED
c.setFont("Helvetica-Bold", 15)
c.drawString(50, 750, "Section 2: Waiting Periods and Limitations")
c.setFont("Helvetica", 11)
c.drawString(50, 720, "2.1 Pre-Existing Diseases (PED) Waiting Period:")
c.drawString(50, 700, "Under this Optima Secure policy, any pre-existing medical condition (PED) disclosed in the")
c.drawString(50, 680, "proposal form has a waiting period of exactly 36 months of continuous coverage.")
c.drawString(50, 650, "2.2 Initial Waiting Period:")
c.drawString(50, 630, "A waiting period of 30 days applies to all illnesses, except accidental injuries which are")
c.drawString(50, 610, "covered from Day 1 without waiting period.")
c.drawString(50, 580, "2.3 Specific Illnesses Waiting Period:")
c.drawString(50, 560, "A 24-month waiting period applies to Cataract, Hernia, Hydrocele, Piles, Sinusitis, Joint")
c.drawString(50, 540, "Replacement surgery, and benign prostatic hypertrophy.")
c.drawString(50, 510, "2.4 Maternity Benefits Waiting Period:")
c.drawString(50, 490, "Maternity and newborn child benefits are not covered under base plan unless the optional")
c.drawString(50, 470, "Maternity Rider is endorsed, which carries a strict 24-month waiting period.")
c.drawString(50, 40, "Page 2 of 6")
c.showPage()

# Page 3: In-Patient Hospitalisation, Room Rent & Sub-limits
c.setFont("Helvetica-Bold", 15)
c.drawString(50, 750, "Section 3: Room Rent, ICU Charges & Hospitalisation Coverage")
c.setFont("Helvetica", 11)
c.drawString(50, 720, "3.1 Room Rent Limit:")
c.drawString(50, 700, "Optima Secure policy has ZERO room rent capping or sub-limits. The policyholder may")
c.drawString(50, 680, "select single private AC room with no proportional deduction on associated medical bills.")
c.drawString(50, 650, "3.2 ICU Charges:")
c.drawString(50, 630, "Intensive Care Unit (ICU) charges and monitoring fees are fully covered up to the Sum Insured.")
c.drawString(50, 600, "3.3 Day Care Procedures:")
c.drawString(50, 580, "All day care procedures requiring less than 24 hours hospitalisation due to advanced medical")
c.drawString(50, 560, "technology are covered up to the full Sum Insured.")
c.drawString(50, 530, "3.4 AYUSH Hospitalisation:")
c.drawString(50, 510, "In-patient hospitalisation under Ayurveda, Yoga, Unani, Siddha, and Homeopathy (AYUSH) is")
c.drawString(50, 490, "covered up to 100% of the Sum Insured in government or Quality Council accredited hospitals.")
c.drawString(50, 40, "Page 3 of 6")
c.showPage()

# Page 4: Value Added Riders, Modern Treatments & Air Ambulance
c.setFont("Helvetica-Bold", 15)
c.drawString(50, 750, "Section 4: Value Added Riders & Special Benefits")
c.setFont("Helvetica", 11)
c.drawString(50, 720, "4.1 Emergency Air Ambulance:")
c.drawString(50, 700, "Air ambulance expenses are covered up to a limit of INR 2,50,000 per hospitalization,")
c.drawString(50, 680, "subject to certified medical necessity for immediate transit to a tertiary hospital.")
c.drawString(50, 650, "4.2 Robotic and Modern Treatments:")
c.drawString(50, 630, "Robotic surgery and stereotactic radio-therapy are covered up to the full Sum Insured,")
c.drawString(50, 610, "subject to a mandatory 10% co-payment payable by the insured.")
c.drawString(50, 580, "4.3 Automatic Restoration Benefit:")
c.drawString(50, 560, "100% automatic restore of basic Sum Insured is triggered once per policy year upon partial")
c.drawString(50, 540, "or full exhaustion, available for subsequent hospitalisations for unrelated illnesses.")
c.drawString(50, 40, "Page 4 of 6")
c.showPage()

# Page 5: Exclusions & Conditional Coverage
c.setFont("Helvetica-Bold", 15)
c.drawString(50, 750, "Section 5: Permanent Exclusions and Conditional Exceptions")
c.setFont("Helvetica", 11)
c.drawString(50, 720, "5.1 Permanent Exclusions:")
c.drawString(50, 700, "Cosmetic and plastic surgery, intentional self-inflicted harm, unproven experimental therapies,")
c.drawString(50, 680, "and adventure sports injuries are permanently excluded from coverage under this policy.")
c.drawString(50, 650, "5.2 Obesity / Bariatric Surgery Exception:")
c.drawString(50, 630, "Bariatric surgery is excluded, EXCEPT when body mass index (BMI) is greater than 40 and")
c.drawString(50, 610, "supported by senior medical specialist certification of severe cardiac or sleep apnea conditions.")
c.drawString(50, 580, "5.3 Dental Treatment:")
c.drawString(50, 560, "Out-patient dental treatments are excluded, UNLESS necessitated by accidental injury")
c.drawString(50, 540, "requiring 24-hour in-patient hospitalisation under general anesthesia.")
c.drawString(50, 40, "Page 5 of 6")
c.showPage()

# Page 6: Claims Procedure, Required Documents & Rejection Grounds
c.setFont("Helvetica-Bold", 15)
c.drawString(50, 750, "Section 6: Claim Process, Required Documents & Rejection Grounds")
c.setFont("Helvetica", 11)
c.drawString(50, 720, "6.1 Cashless Claim Notification:")
c.drawString(50, 700, "Planned hospital admission must be pre-authorised 48 hours prior to admission.")
c.drawString(50, 680, "Emergency hospitalisation must be intimated within 24 hours of hospital admission.")
c.drawString(50, 650, "6.2 Reimbursement Claim Submission Window:")
c.drawString(50, 630, "All original hospital bills, payment receipts, discharge summary, diagnostic reports, and")
c.drawString(50, 610, "pharmacy invoices must be submitted within 30 days from discharge date.")
c.drawString(50, 580, "6.3 Grounds for Claim Rejection:")
c.drawString(50, 560, "Claims will be rejected for: (a) Fraudulent or inflated bills, (b) Non-disclosure of material")
c.drawString(50, 540, "pre-existing conditions, (c) Treatments during applicable waiting periods, and (d) Submission")
c.drawString(50, 520, "after 30 days without medical justification.")
c.drawString(50, 40, "Page 6 of 6")
c.showPage()

c.save()
print("Generated Comprehensive PDF:", pdf_path)

# ---------------------------------------------------------
# 2. Generate Sample DOCX with Headings & Markdown Tables
# ---------------------------------------------------------
doc_path = "tests/sample_docs/Star_Health_Policy_Schedule.docx"
doc = docx.Document()
doc.add_heading("Star Health Allied Insurance - Comprehensive Schedule", level=1)
doc.add_paragraph("Policy Name: Star Health Comprehensive Insurance Policy")
doc.add_paragraph("Policy Number: STAR-COMP-2026-443210")
doc.add_paragraph("Pre-Existing Condition Waiting Period: 36 months from initial policy inception.")
doc.add_paragraph("Room Rent Category: Single standard AC room with no capping.")
doc.add_paragraph("Maternity Benefit: Normal delivery covered up to INR 25,000 and Caesarean up to INR 40,000 after 24 months continuous coverage.")
doc.add_paragraph("Newborn Baby Cover: Covered up to INR 1,00,000 from Day 1 subject to valid maternity claim.")

# Table of Benefit Limits
table = doc.add_table(rows=1, cols=3)
hdr_cells = table.rows[0].cells
hdr_cells[0].text = "Benefit Description"
hdr_cells[1].text = "Limit / Sub-limit"
hdr_cells[2].text = "Waiting Period"

benefits = [
    ("Pre-Existing Diseases", "Full Sum Insured", "36 Months"),
    ("Cataract Surgery", "INR 30,000 per eye", "24 Months"),
    ("Maternity (Normal)", "INR 25,000", "24 Months"),
    ("Maternity (Caesarean)", "INR 40,000", "24 Months"),
    ("Health Checkup", "INR 2,500 per year", "1 Year Claim-free")
]
for desc, lim, wp in benefits:
    row_cells = table.add_row().cells
    row_cells[0].text = desc
    row_cells[1].text = lim
    row_cells[2].text = wp

doc.save(doc_path)
print("Generated Comprehensive DOCX:", doc_path)

# ---------------------------------------------------------
# 3. Generate Sample TXT with Headings and Page Dividers
# ---------------------------------------------------------
txt_path = "tests/sample_docs/Care_Health_Policy.txt"
with open(txt_path, "w", encoding="utf-8") as f:
    f.write("--- Page 1 ---\n")
    f.write("Care Health Insurance Policy Summary\n")
    f.write("Policy Number: CARE-SUPREME-2026-990011\n")
    f.write("Sum Insured: INR 15,00,000\n\n")
    f.write("Section 1: Waiting Periods\n")
    f.write("Pre-existing disease waiting period is strictly 36 months.\n")
    f.write("Initial waiting period is 30 days for all non-accidental hospitalisations.\n\n")
    f.write("--- Page 2 ---\n")
    f.write("Section 2: Bonus & Riders\n")
    f.write("No Claim Bonus super add-on increases sum insured by 50% each claim-free year up to 500%.\n")
    f.write("Emergency international coverage up to USD 50,000 for life-threatening emergencies.\n")
    f.write("OPD consultations capped at INR 5,000 per policy year.\n")
print("Generated Comprehensive TXT:", txt_path)

# ---------------------------------------------------------
# 4. Generate Blank/Scanned PDF for Negative OCR Testing
# ---------------------------------------------------------
scanned_pdf_path = "tests/sample_docs/Scanned_Sample_Policy.pdf"
sc = canvas.Canvas(scanned_pdf_path, pagesize=letter)
# Add empty blank page with no digital text layer
sc.showPage()
sc.save()
print("Generated Scanned/Blank Test PDF:", scanned_pdf_path)
